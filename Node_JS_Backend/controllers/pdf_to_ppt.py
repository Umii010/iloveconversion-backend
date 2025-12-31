import sys
import os
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches

pdf_path = sys.argv[1]
ppt_path = sys.argv[2]

images = convert_from_path(pdf_path, dpi=200)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

for img in images:
    slide = prs.slides.add_slide(blank_slide_layout)

    img_path = os.path.join(os.path.dirname(ppt_path), "temp_page.png")
    img.save(img_path, "PNG")

    slide.shapes.add_picture(
        img_path,
        Inches(0),
        Inches(0),
        width=prs.slide_width,
        height=prs.slide_height
    )

    os.remove(img_path)

prs.save(ppt_path)
